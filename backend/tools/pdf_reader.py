import os
from pypdf import PdfReader

try:
    import docx
except ImportError:
    docx = None

try:
    import pptx
except ImportError:
    pptx = None


def read_pdf(

    path

):
    _, ext = os.path.splitext(path.lower())

    if ext == ".pdf":
        reader = PdfReader(

            path

        )

        text = ""

        for page in reader.pages:

            page_text = page.extract_text()

            if page_text:

                text += page_text


        return text

    elif ext in [".docx", ".doc"]:
        if docx is None:
            return "Word parsing library python-docx is not installed."
        try:
            doc = docx.Document(path)
            text = []
            for para in doc.paragraphs:
                if para.text:
                    text.append(para.text)
            for table in doc.tables:
                for row in table.rows:
                    row_text = [cell.text for cell in row.cells if cell.text]
                    if row_text:
                        text.append(" | ".join(row_text))
            return "\n".join(text)
        except Exception as e:
            return f"Error reading Word document: {str(e)}"

    elif ext in [".pptx", ".ppt"]:
        if pptx is None:
            return "PowerPoint parsing library python-pptx is not installed."
        try:
            prs = pptx.Presentation(path)
            text = []
            for slide in prs.slides:
                slide_text = []
                for shape in slide.shapes:
                    if hasattr(shape, "text") and shape.text:
                        slide_text.append(shape.text)
                if slide_text:
                    text.append("\n".join(slide_text))
            return "\n---\n".join(text)
        except Exception as e:
            return f"Error reading PowerPoint presentation: {str(e)}"

    elif ext in [".txt", ".md", ".json", ".csv", ".tsv"]:
        try:
            with open(path, "r", encoding="utf-8", errors="ignore") as f:
                return f.read()
        except Exception as e:
            return f"Error reading text file: {str(e)}"

    else:
        try:
            with open(path, "r", encoding="utf-8", errors="ignore") as f:
                return f.read()
        except Exception:
            return f"Unsupported file format: {ext}"